import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from openpyxl import load_workbook

MAX_ROWS = 5000

ALIASES = {
    "sku": {"sku", "codigo", "codigo_produto", "cod_produto", "referencia", "ref", "id_produto"},
    "gtin": {"gtin", "ean", "ean13", "codigo_barras", "codigo_de_barras"},
    "name": {"nome", "produto", "nome_produto", "descricao", "titulo", "title"},
    "category": {"categoria", "category", "departamento"},
    "cost": {"custo", "preco_custo", "preco_atacado", "valor_atacado", "cost", "preco_compra"},
    "stock": {"estoque", "stock", "quantidade", "qtd", "saldo"},
    "sale_price": {"preco_venda", "preco_varejo", "valor_varejo", "varejo", "sale_price", "preco_sugerido"},
    "shipping_hours": {"prazo_horas", "shipping_hours", "prazo_envio_horas", "horas_envio"},
}

REQUIRED = {"sku", "name", "cost", "stock", "sale_price"}

@dataclass
class ParsedSupplierRow:
    sku: str
    gtin: str | None
    name: str
    category: str
    cost: float
    stock: int
    sale_price: float
    shipping_hours: int

def normalize_header(value: object) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")

def canonical_header(value: object) -> str | None:
    norm = normalize_header(value)
    for key, aliases in ALIASES.items():
        if norm in aliases:
            return key
    return None

def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Não foi possível identificar a codificação do CSV.")

def _csv_rows(content: bytes) -> list[dict]:
    text = _decode_csv(content)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    return list(reader)

def _xlsx_rows(content: bytes) -> list[dict]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(v or "") for v in next(rows)]
    except StopIteration:
        return []
    result = []
    for values in rows:
        result.append({headers[i]: values[i] if i < len(values) else None for i in range(len(headers))})
    return result

def _to_float(value: object) -> float:
    if value is None or str(value).strip() == "":
        raise ValueError("valor vazio")
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("R$", "").replace(" ", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(".", "").replace(",", ".")
    return float(text)

def _to_int(value: object) -> int:
    return max(0, int(round(_to_float(value))))

def parse_supplier_file(filename: str, content: bytes) -> tuple[list[ParsedSupplierRow], list[str], dict]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        raw_rows = _csv_rows(content)
    elif lower.endswith(".xlsx"):
        raw_rows = _xlsx_rows(content)
    else:
        raise ValueError("Formato não suportado. Envie um arquivo .csv ou .xlsx.")

    if not raw_rows:
        raise ValueError("O arquivo está vazio.")

    raw_headers = list(raw_rows[0].keys())
    mapping = {}
    for header in raw_headers:
        canonical = canonical_header(header)
        if canonical and canonical not in mapping:
            mapping[canonical] = header

    missing = sorted(REQUIRED - set(mapping))
    if missing:
        pretty = {
            "sku": "SKU/código",
            "name": "produto/nome",
            "cost": "custo/preço atacado",
            "stock": "estoque",
            "sale_price": "preço de venda/varejo",
        }
        raise ValueError("Colunas obrigatórias não encontradas: " + ", ".join(pretty[x] for x in missing))

    parsed: list[ParsedSupplierRow] = []
    errors: list[str] = []

    for index, raw in enumerate(raw_rows[:MAX_ROWS], start=2):
        try:
            sku = str(raw.get(mapping["sku"], "") or "").strip()
            name = str(raw.get(mapping["name"], "") or "").strip()
            if not sku or not name:
                raise ValueError("SKU e nome são obrigatórios")
            cost = _to_float(raw.get(mapping["cost"]))
            stock = _to_int(raw.get(mapping["stock"]))
            sale_price = _to_float(raw.get(mapping["sale_price"]))
            if cost < 0 or sale_price <= 0:
                raise ValueError("preços inválidos")
            gtin = str(raw.get(mapping.get("gtin"), "") or "").strip() or None
            category = str(raw.get(mapping.get("category"), "") or "Geral").strip() or "Geral"
            shipping_hours = 24
            if mapping.get("shipping_hours") and raw.get(mapping["shipping_hours"]) not in (None, ""):
                shipping_hours = max(1, _to_int(raw.get(mapping["shipping_hours"])))
            parsed.append(ParsedSupplierRow(
                sku=sku,
                gtin=gtin,
                name=name,
                category=category,
                cost=round(cost, 2),
                stock=stock,
                sale_price=round(sale_price, 2),
                shipping_hours=shipping_hours,
            ))
        except Exception as exc:
            errors.append(f"Linha {index}: {exc}")

    meta = {
        "rows_read": min(len(raw_rows), MAX_ROWS),
        "rows_accepted": len(parsed),
        "rows_rejected": len(errors),
        "detected_columns": {key: str(value) for key, value in mapping.items()},
        "truncated": len(raw_rows) > MAX_ROWS,
    }
    return parsed, errors, meta
